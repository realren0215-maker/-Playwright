from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import time
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, sync_playwright


@dataclass
class StoreAccount:
    store_name: str
    username: str
    password: str


@dataclass
class StoreResult:
    store_name: str
    status: str
    downloaded_file: str = ""
    reason: str = ""


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        config = yaml.safe_load(f)
    if not isinstance(config, dict):
        raise ValueError("config.yaml 형식이 올바르지 않습니다.")
    return config


def normalize_account_id(value: Any) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_accounts(config: dict[str, Any], stores_filter: list[str] | None, limit: int | None) -> list[StoreAccount]:
    columns = config["account_columns"]
    df = pd.read_excel(config["account_excel"], dtype={columns["username"]: object, columns["password"]: object})
    required = [columns["store"], columns["username"], columns["password"]]
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError(f"계정 엑셀에 필요한 컬럼이 없습니다: {missing}")

    accounts: list[StoreAccount] = []
    wanted = {s.strip() for s in stores_filter} if stores_filter else None
    for _, row in df.iterrows():
        store_name = str(row[columns["store"]]).strip()
        if not store_name or store_name == "nan":
            continue
        if wanted and store_name not in wanted:
            continue
        accounts.append(
            StoreAccount(
                store_name=store_name,
                username=normalize_account_id(row[columns["username"]]),
                password=str(row[columns["password"]]).strip(),
            )
        )
    return accounts[:limit] if limit else accounts


def get_run_date(config: dict[str, Any]) -> date:
    configured = config.get("date_range", {}).get("run_date")
    if configured:
        return datetime.strptime(str(configured), "%Y-%m-%d").date()
    return date.today()


def get_query_range(config: dict[str, Any]) -> tuple[date, date, date]:
    run_date = get_run_date(config)
    start = run_date.replace(day=1)
    end = run_date - timedelta(days=1)
    if end < start:
        raise ValueError(
            f"조회 기간을 만들 수 없습니다. 실행일 {run_date:%Y-%m-%d} 기준 종료일이 월 1일보다 빠릅니다."
        )
    return run_date, start, end


def safe_sheet_title(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", name).strip() or "Sheet"
    base = cleaned[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f"_{counter}"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate)
    return candidate


def safe_filename(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\n\r\t]+', "_", name).strip() or "store"


def click_first_visible_text(page: Page, text: str, timeout_ms: int) -> bool:
    candidates = [
        page.get_by_text(text, exact=True),
        page.get_by_text(text),
        page.locator(f"text={text}"),
    ]
    for locator in candidates:
        try:
            locator.first.click(timeout=timeout_ms)
            return True
        except PlaywrightTimeoutError:
            continue
        except Exception:
            continue
    return False


def fill_by_placeholder(page: Page, placeholder: str, value: str, timeout_ms: int) -> None:
    field = page.get_by_placeholder(placeholder).first
    field.fill(value, timeout=timeout_ms)


def wait_for_possible_manual_verification(page: Page, config: dict[str, Any]) -> None:
    keywords = config["login"].get("verification_keywords", [])
    timeout = int(config.get("manual_verification_timeout_sec", 180))
    detected = False
    for keyword in keywords:
        try:
            if page.get_by_text(keyword).first.is_visible(timeout=1500):
                detected = True
                break
        except Exception:
            continue
    if not detected:
        return

    print("\n[인증 필요] 브라우저에서 슬라이드/인증번호 처리를 완료한 뒤 Enter를 누르세요.")
    print(f"{timeout}초 안에 처리하지 못하면 실패로 기록됩니다.")
    start = time.time()
    input("인증 완료 후 Enter: ")
    if time.time() - start > timeout:
        raise TimeoutError("수동 인증 대기 시간이 초과되었습니다.")


def login(page: Page, account: StoreAccount, config: dict[str, Any]) -> None:
    timeout_ms = int(config.get("timeout_ms", 30000))
    page.goto(config["login_url"], wait_until="domcontentloaded", timeout=timeout_ms)
    page.wait_for_load_state("networkidle", timeout=timeout_ms)

    click_first_visible_text(page, config["login"]["account_tab_text"], timeout_ms=5000)
    fill_by_placeholder(page, config["login"]["username_placeholder"], account.username, timeout_ms)
    fill_by_placeholder(page, config["login"]["password_placeholder"], account.password, timeout_ms)

    for checkbox in [
        page.locator("input[type='checkbox']").first,
        page.locator(".el-checkbox__input").first,
    ]:
        try:
            checkbox.check(force=True, timeout=3000)
            break
        except Exception:
            try:
                checkbox.click(force=True, timeout=3000)
                break
            except Exception:
                continue

    if not click_first_visible_text(page, config["login"]["login_button_text"], timeout_ms):
        raise RuntimeError("로그인 버튼을 찾지 못했습니다.")

    wait_for_possible_manual_verification(page, config)
    page.wait_for_load_state("networkidle", timeout=timeout_ms)


def navigate_to_report(page: Page, config: dict[str, Any]) -> None:
    timeout_ms = int(config.get("timeout_ms", 30000))
    clicked_any = False
    for text in config["navigation"].get("click_texts", []):
        if click_first_visible_text(page, text, timeout_ms=8000):
            clicked_any = True
            page.wait_for_timeout(800)

    if not clicked_any and config["navigation"].get("fallback_report_url"):
        page.goto(config["navigation"]["fallback_report_url"], wait_until="networkidle", timeout=timeout_ms)

    # 데이터 보고 탭이 이미 보이면 한 번 더 눌러 확실히 진입합니다.
    click_first_visible_text(page, "数据报告", timeout_ms=5000)
    page.wait_for_load_state("networkidle", timeout=timeout_ms)


def set_date_range(page: Page, config: dict[str, Any], start: date, end: date) -> None:
    timeout_ms = int(config.get("timeout_ms", 30000))
    flow = config["download_flow"]
    click_first_visible_text(page, flow.get("custom_date_text", "自定义"), timeout_ms=8000)
    page.wait_for_timeout(500)

    start_text = start.strftime("%Y-%m-%d")
    end_text = end.strftime("%Y-%m-%d")
    candidates = [
        "input[placeholder*='开始']",
        "input[placeholder*='开始日期']",
        "input[placeholder*='请选择开始']",
    ]
    end_candidates = [
        "input[placeholder*='结束']",
        "input[placeholder*='结束日期']",
        "input[placeholder*='请选择结束']",
    ]

    filled = False
    for start_selector in candidates:
        try:
            page.locator(start_selector).first.fill(start_text, timeout=3000)
            filled = True
            break
        except Exception:
            continue
    for end_selector in end_candidates:
        try:
            page.locator(end_selector).first.fill(end_text, timeout=3000)
            filled = True
            break
        except Exception:
            continue

    if not filled:
        # Element UI의 range input은 placeholder가 없을 수 있어 보이는 입력창 뒤쪽부터 시도합니다.
        inputs = page.locator("input:visible")
        count = inputs.count()
        for idx in range(max(0, count - 4), count):
            try:
                value = start_text if idx == count - 2 else end_text if idx == count - 1 else None
                if value:
                    inputs.nth(idx).fill(value, timeout=3000)
                    filled = True
            except Exception:
                continue

    if not filled:
        raise RuntimeError("조회 기간 입력창을 찾지 못했습니다. config.yaml의 날짜 관련 selector 조정이 필요합니다.")

    page.keyboard.press("Enter")
    page.wait_for_load_state("networkidle", timeout=timeout_ms)


def open_download_record(page: Page, config: dict[str, Any]) -> None:
    flow = config["download_flow"]
    timeout_ms = int(config.get("timeout_ms", 30000))
    if click_first_visible_text(page, flow.get("popup_go_text", "前往获取"), timeout_ms=8000):
        page.wait_for_timeout(1000)
        return
    click_first_visible_text(page, flow.get("download_record_text", "下载记录"), timeout_ms=8000)
    page.wait_for_timeout(1000)


def download_store_file(page: Page, account: StoreAccount, config: dict[str, Any], raw_dir: Path) -> Path:
    timeout_ms = int(config.get("timeout_ms", 30000))
    flow = config["download_flow"]

    if not click_first_visible_text(page, flow.get("download_detail_text", "下载明细"), timeout_ms):
        raise RuntimeError("데이터 다운로드 버튼(下载明细)을 찾지 못했습니다.")

    open_download_record(page, config)

    deadline = time.time() + int(flow.get("poll_download_record_sec", 90))
    while time.time() < deadline:
        try:
            complete_visible = page.get_by_text(flow.get("complete_text", "已完成")).first.is_visible(timeout=3000)
        except Exception:
            complete_visible = False
        if complete_visible:
            break
        page.wait_for_timeout(3000)
    else:
        raise TimeoutError("다운로드 기록에서 완료 상태를 확인하지 못했습니다.")

    with page.expect_download(timeout=timeout_ms) as download_info:
        if not click_first_visible_text(page, flow.get("row_download_text", "下载"), timeout_ms):
            raise RuntimeError("다운로드 기록의 다운로드 링크를 찾지 못했습니다.")
    download = download_info.value
    suggested_name = download.suggested_filename or f"{safe_filename(account.store_name)}.xlsx"
    suffix = Path(suggested_name).suffix or ".xlsx"
    destination = raw_dir / f"{safe_filename(account.store_name)}{suffix}"
    download.save_as(destination)
    return destination


def copy_worksheet(source_path: Path, target_wb: Workbook, sheet_name: str) -> None:
    source_wb = load_workbook(source_path)
    source_ws = source_wb[source_wb.sheetnames[0]]
    target_ws = target_wb.create_sheet(sheet_name)

    for row in source_ws.iter_rows():
        for cell in row:
            new_cell = target_ws[cell.coordinate]
            new_cell.value = cell.value
            if cell.has_style:
                new_cell.font = copy(cell.font)
                new_cell.fill = copy(cell.fill)
                new_cell.border = copy(cell.border)
                new_cell.alignment = copy(cell.alignment)
                new_cell.number_format = cell.number_format
                new_cell.protection = copy(cell.protection)
            if cell.hyperlink:
                new_cell._hyperlink = copy(cell.hyperlink)
            if cell.comment:
                new_cell.comment = copy(cell.comment)

    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))
    for col_idx, dim in source_ws.column_dimensions.items():
        target_ws.column_dimensions[col_idx].width = dim.width
    for row_idx, dim in source_ws.row_dimensions.items():
        target_ws.row_dimensions[row_idx].height = dim.height
    target_ws.freeze_panes = source_ws.freeze_panes
    if source_ws.auto_filter.ref:
        target_ws.auto_filter.ref = source_ws.auto_filter.ref

    for idx in range(1, source_ws.max_column + 1):
        col = get_column_letter(idx)
        if target_ws.column_dimensions[col].width is None:
            target_ws.column_dimensions[col].width = 14


def merge_downloads(results: list[StoreResult], output_file: Path) -> None:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    used: set[str] = set()

    ok_results = [r for r in results if r.status == "success" and r.downloaded_file]
    if not ok_results:
        raise RuntimeError("통합할 다운로드 파일이 없습니다.")

    for result in ok_results:
        sheet_name = safe_sheet_title(result.store_name, used)
        copy_worksheet(Path(result.downloaded_file), wb, sheet_name)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def write_failure_log(results: list[StoreResult], path: Path) -> None:
    failures = [r for r in results if r.status != "success"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["매장명", "상태", "실패사유", "다운로드파일"])
        writer.writeheader()
        for r in failures:
            writer.writerow(
                {
                    "매장명": r.store_name,
                    "상태": r.status,
                    "실패사유": r.reason,
                    "다운로드파일": r.downloaded_file,
                }
            )


def run_downloads(config: dict[str, Any], accounts: list[StoreAccount], raw_dir: Path, start: date, end: date) -> list[StoreResult]:
    results: list[StoreResult] = []
    timeout_ms = int(config.get("timeout_ms", 30000))

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=bool(config.get("headless", False)),
            slow_mo=int(config.get("slow_mo_ms", 0)),
        )

        for idx, account in enumerate(accounts, start=1):
            print(f"\n[{idx}/{len(accounts)}] {account.store_name} 처리 시작")
            context = browser.new_context(accept_downloads=True, locale="zh-CN")
            page = context.new_page()
            page.set_default_timeout(timeout_ms)
            try:
                login(page, account, config)
                navigate_to_report(page, config)
                set_date_range(page, config, start, end)
                downloaded = download_store_file(page, account, config, raw_dir)
                results.append(StoreResult(account.store_name, "success", str(downloaded)))
                print(f"[성공] {account.store_name}: {downloaded}")
            except Exception as exc:
                screenshot = raw_dir / f"{safe_filename(account.store_name)}_failure.png"
                try:
                    page.screenshot(path=screenshot, full_page=True)
                except Exception:
                    pass
                reason = f"{type(exc).__name__}: {exc}"
                results.append(StoreResult(account.store_name, "failed", reason=reason))
                print(f"[실패] {account.store_name}: {reason}")
            finally:
                page.close()
                context.close()
        browser.close()

    return results


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="따중디엔핑 백오피스 주간데이터 자동 다운로드/통합")
    parser.add_argument("--config", default="config.yaml", help="config.yaml 경로")
    parser.add_argument("--limit", type=int, default=None, help="테스트 매장 수. config.yaml보다 우선합니다.")
    parser.add_argument("--stores", default=None, help="쉼표로 구분한 특정 매장명 목록")
    parser.add_argument("--headless", action="store_true", help="브라우저를 숨기고 실행")
    parser.add_argument("--merge-only", action="store_true", help="이미 다운로드된 raw 폴더 파일만 통합")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config_path = Path(args.config).resolve()
    config = load_config(config_path)
    if args.headless:
        config["headless"] = True

    run_date, start, end = get_query_range(config)
    output_dir = Path(config["output_dir"])
    if not output_dir.is_absolute():
        output_dir = config_path.parent / output_dir
    raw_dir = output_dir / "raw" / run_date.strftime("%Y-%m-%d")
    raw_dir.mkdir(parents=True, exist_ok=True)

    final_file = output_dir / f"주간데이터_{run_date:%Y-%m-%d}.xlsx"
    failure_log = output_dir / f"failure_log_{run_date:%Y-%m-%d}.csv"

    stores_filter = [s.strip() for s in args.stores.split(",")] if args.stores else None
    limit = args.limit if args.limit is not None else config.get("test_limit")
    accounts = load_accounts(config, stores_filter, limit)
    if not accounts and not args.merge_only:
        raise RuntimeError("실행할 매장 계정이 없습니다.")

    print(f"조회 기간: {start:%Y-%m-%d} ~ {end:%Y-%m-%d}")
    print(f"최종 파일: {final_file}")

    if args.merge_only:
        results = [
            StoreResult(path.stem, "success", str(path))
            for path in sorted(raw_dir.glob("*.xls*"))
            if not path.name.startswith("~$")
        ]
    else:
        results = run_downloads(config, accounts, raw_dir, start, end)

    write_failure_log(results, failure_log)
    success_count = sum(1 for r in results if r.status == "success")
    if success_count:
        merge_downloads(results, final_file)
        print(f"\n통합 완료: {final_file}")
    else:
        print("\n성공한 다운로드가 없어 통합 파일을 만들지 않았습니다.")

    failures = [r for r in results if r.status != "success"]
    if failures:
        print(f"실패 로그: {failure_log}")
    return 0 if success_count else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n사용자 중단")
        raise SystemExit(130)
